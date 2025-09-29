import re
import openpyxl


def find_sheets(excel_file):
    for name in excel_file.wb.sheetnames:
        if re.search(r"calculation.*results", name, re.IGNORECASE):
            excel_file.calculation_sheet = name
        elif re.search(r"payment.*transactions?", name, re.IGNORECASE):
            excel_file.payment_sheet = name

    if not excel_file.calculation_sheet:
        raise ValueError("Calculation Results sheet not found.")
    if not excel_file.payment_sheet:
        raise ValueError("Payment Transactions sheet not found.")


def find_column_to_sum(excel_file):
    sheet = excel_file.wb[excel_file.payment_sheet]
    for cell in sheet[1]:
        if cell.value and "euroamount" in str(cell.value).lower().replace(" ", ""):
            excel_file.column_to_sum = cell.column_letter
            return
    raise ValueError("EuroAmount column not found.")


def extract_claim_amount(excel_file):
    sheet = excel_file.wb[excel_file.calculation_sheet]
    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=2):
        if isinstance(row[0].value, str) and "claim amount" in row[0].value.lower():
            try:
                return float(str(row[1].value).replace(",", ""))
            except:
                continue
    raise ValueError("Claim amount not found.")


def find_transaction_sheets(xlsx_path):
    """
    Heuristic: sheets named like 'All', 'Payment Transactions', 'Receipts', or 'Invoices' are considered transaction sheets.
    """
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    transaction_sheets = [
        sheet.title
        for sheet in wb.worksheets
        if re.search(
            r"^(All|Payment Transactions|Receipts|Invoices)$",
            sheet.title,
            re.IGNORECASE,
        )
    ]
    wb.close()
    return transaction_sheets

import os
import re
from datetime import datetime
import customtkinter as ctk
import tkinter as tk
from tkinter import messagebox
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill


def create_blank_payment_summary_sheet(wb, sheet_name):
    """Creates an empty Payment Summary sheet with standard headers."""
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    headers = ["Date", "Type", "Currency", "Amount", "Status"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c).value = h
        ws.cell(row=1, column=c).font = Font(bold=True)
    return ws


from collections import defaultdict
from processor.utils import (
    try_parse_datetime,
    is_failed_status,
    excel_fmt_date,
    excel_fmt_number,
    parse_any_datetime,
    is_success_status,
    format_summary_sheet,
    safe_float,
    normalize_currency,
)
from processor.normalization import (
    normalize_tx_type,
    normalize_status,
    parse_date,
    parse_amount,
    is_effective_deposit,
    is_effective_withdraw,
    is_rejected,
)


# --- Header detection helper ---
def _find_header_row(sheet, log):
    """
    Robust PokerStars header detection:
    - Scan entire sheet
    - Accept row if it contains 'Amount' (or 'Billing Amount') + any 2 of expected tokens
    - If a 'User ID' banner exists, search starts *after* it
    - Final fallback: row that contains both 'TransId' and ('Amount' or 'Billing Amount')
    """
    import re
    from datetime import datetime

    def norm(s):
        return re.sub(r"\W+", "", str(s or "").strip().lower())

    expected = {
        "transid",
        "parentid",
        "transtype",
        "type",  # NEW
        "started",
        "completed",
        "accountcurrency",
        "currency",  # NEW
        "billingcurrency",  # NEW
        "amount",
        "billingamount",  # NEW
        "status",
        "paysystem",
        "country",
        "loggedin",
        "reference",
        "gateway",  # NEW (matches GateWay)
        "descriptor",  # NEW
        "originaltrans",  # NEW
    }

    # If there is a "User ID" banner, remember last such row to start search after it
    start_after = 1
    for r in range(1, sheet.max_row + 1):
        vals = [str(c.value or "").strip() for c in sheet[r]]
        if any(isinstance(v, str) and v.lower().startswith("user id") for v in vals):
            start_after = r + 1

    def row_is_header(idx):
        raw = [c.value for c in sheet[idx]]
        vals = [str(v or "").strip() for v in raw]
        toks = [norm(v) for v in vals]
        if not any(toks):  # empty row
            return False, vals

        # accept either "amount" OR "billingamount"
        has_amount = ("amount" in toks) or ("billingamount" in toks)
        present = [t for t in toks if t in expected]

        # 'Amount' (or BillingAmount) + any two others (permissive but effective)
        if has_amount and len(set(present)) >= 3:
            return True, vals

        return False, vals

    # Pass 1: strict(ish) pass
    for r in range(start_after, sheet.max_row + 1):
        ok, vals = row_is_header(r)
        if ok:
            log(f"✅ Detected header row at Excel row {r}: {vals}")
            return r, vals

    # Pass 2: fallback — look for 'TransId' & ('Amount' or 'BillingAmount') in same row
    for r in range(start_after, sheet.max_row + 1):
        vals = [str(c.value or "").strip() for c in sheet[r]]
        toks = [norm(v) for v in vals]
        if "transid" in toks and ("amount" in toks or "billingamount" in toks):
            log(f"✅ Detected header row (fallback) at Excel row {r}: {vals}")
            return r, vals

    log(f"⚠️ No distinct header row found (scanned {sheet.max_row} rows).")
    return None, []


POSITIVE_OK = {
    "processed",
    "completed",
    "approved",
    "complete",
    "dep_settled",
    "pay_closed",
    "closed",
    "settled",
    "posted",
}
NEGATIVE_BAD = {"rejected", "failed", "declined", "pending", "cancelled", "canceled"}

WITHDRAW_KEYWORDS = {"redeem", "withdraw", "withdrawal", "cashout", "payout"}
DEPOSIT_KEYWORDS = {"deposit", "storting"}


def create_summary_sheet(excel_file, parsed_rows, log):
    wb = excel_file.wb
    if "Summary" in wb.sheetnames:
        del wb["Summary"]

    summary_sheet = wb.create_sheet("Summary")

    if "Payment Summary" not in wb.sheetnames:
        raise ValueError("Missing 'Payment Summary' sheet to generate accurate totals.")

    pws = wb["Payment Summary"]
    header = [cell.value for cell in pws[1]]

    # Column indexes (Currency optional)
    date_idx = header.index("Date")
    type_idx = header.index("Type")
    amount_idx = header.index("Amount")
    currency_idx = header.index("Currency") if "Currency" in header else None

    # Collect data: currency -> type -> year -> amount
    data = defaultdict(lambda: defaultdict(lambda: defaultdict(float)))
    all_years, all_dates, all_currencies = set(), [], set()

    for row in pws.iter_rows(min_row=2, values_only=True):
        try:
            dt = row[date_idx]
            tx_type = row[type_idx]
            amt = row[amount_idx]
            cur = row[currency_idx] if currency_idx is not None else "—"

            if not isinstance(dt, (str, int)) and hasattr(dt, "year"):
                yr = dt.year
                data[cur][tx_type][yr] += amt
                all_years.add(yr)
                all_dates.append(dt)
                all_currencies.add(cur)
        except Exception:
            continue

    if not all_dates:
        log("⚠️ No valid dates found in Payment Summary.")
        return

    start_date = min(all_dates).strftime("%d.%m.%Y")
    end_date = max(all_dates).strftime("%d.%m.%Y")
    years = sorted(all_years)
    currencies = sorted(all_currencies, key=lambda c: (c != "—", str(c)))

    # Title
    title = f"Summary Payment Transactions ({start_date} - {end_date})"
    summary_sheet.merge_cells(
        start_row=1, start_column=1, end_row=1, end_column=len(years) + 3
    )
    title_cell = summary_sheet.cell(row=1, column=1, value=title)
    title_cell.font = Font(size=14, bold=True)
    title_cell.alignment = Alignment(horizontal="center")

    # Headers
    headers = (
        ["Currency", "Transaction Type"] + [str(y) for y in years] + ["Grand Total"]
    )
    summary_sheet.append(headers)

    # Rows (Currency → Type)
    tx_types = ["Deposit", "Redeem"]
    for cur in currencies:
        for tx in tx_types:
            row = [cur, tx]
            total = 0.0
            for y in years:
                val = round(data[cur][tx].get(y, 0.0), 2)
                row.append(val)
                total += val
            row.append(round(total, 2))
            summary_sheet.append(row)

    # Grand total row
    grand = ["Grand Total", ""]
    for i in range(2, len(years) + 3):  # sum each numeric column
        col_letter = get_column_letter(i + 1)  # shift because first column is A
        start_row = 3
        end_row = 2 + len(currencies) * len(tx_types)
        grand.append(f"=SUM({col_letter}{start_row}:{col_letter}{end_row})")
    summary_sheet.append(grand)

    # Table
    end_col_letter = get_column_letter(len(headers))
    end_row = 2 + len(currencies) * len(tx_types) + 1
    table_range = f"A2:{end_col_letter}{end_row}"
    table = Table(displayName="SummaryTable", ref=table_range)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9", showRowStripes=True, showColumnStripes=False
    )
    summary_sheet.add_table(table)

    # Styling
    header_fill = PatternFill(
        start_color="4472C4", end_color="4472C4", fill_type="solid"
    )
    white_font = Font(color="FFFFFF", bold=True)

    for cell in summary_sheet[2]:
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center")

    # Number formatting
    for row in summary_sheet.iter_rows(min_row=3, max_row=end_row):
        for cell in row[2:]:  # numeric cols only
            try:
                if isinstance(cell.value, (int, float)) or (
                    isinstance(cell.value, str) and cell.value.startswith("=")
                ):
                    cell.number_format = "#,##0.00"
            except:
                pass

    # Auto-width
    for col in summary_sheet.columns:
        max_len = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        summary_sheet.column_dimensions[get_column_letter(col[0].column)].width = (
            max_len + 2
        )

    log("✅ Created Summary sheet (Currency → Type → Year).")


def _select_multiple_sheets(wb, title="Select Payment History Sheet"):
    """
    Modal multi-select for worksheet names. Returns a list of selected sheet names.
    If user cancels/closes with nothing selected, returns [].
    """
    selected = []

    # Create a transient modal window
    win = ctk.CTkToplevel()
    win.title(title)
    win.geometry("420x420")
    win.grab_set()  # modal
    win.focus_force()
    try:
        # Make it feel modal even without a known parent
        win.attributes("-topmost", True)
    except Exception:
        pass

    header = ctk.CTkLabel(win, text="Tick all sheets you want to include:", anchor="w")
    header.pack(padx=12, pady=10, fill="x")

    scroll = ctk.CTkScrollableFrame(win)
    scroll.pack(padx=12, pady=(0, 12), fill="both", expand=True)

    vars_by_name = {}
    for name in wb.sheetnames:
        v = tk.BooleanVar(value=False)
        chk = ctk.CTkCheckBox(scroll, text=name, variable=v)
        chk.pack(anchor="w", padx=8, pady=4)
        vars_by_name[name] = v

    btn_row = ctk.CTkFrame(win)
    btn_row.pack(pady=10, fill="x")

    def on_ok():
        for n, v in vars_by_name.items():
            if v.get():
                selected.append(n)
        win.destroy()

    def on_sel_all():
        for v in vars_by_name.values():
            v.set(True)

    def on_clear():
        for v in vars_by_name.values():
            v.set(False)

    sel_all_btn = ctk.CTkButton(btn_row, text="Select all", command=on_sel_all)
    sel_all_btn.pack(side="left", padx=6)

    clear_btn = ctk.CTkButton(btn_row, text="Clear", command=on_clear)
    clear_btn.pack(side="left", padx=6)

    ok_btn = ctk.CTkButton(btn_row, text="OK", command=on_ok)
    ok_btn.pack(side="right", padx=6)

    win.wait_window()  # block until closed
    return selected


def create_ta_payment_summary(excel_file, log):
    """
    Builds 'Payment Summary' from:
      - a single combined sheet, OR
      - user-selected Deposit/Cashout sheets.

    Strategy:
      • If user picked the SAME sheet on both sides → parse once with the combined extractor.
      • If user picked distinct Deposit/Cashout sheets → use the classic per-sheet extractor
        (this is what worked for Bwin/PokerStars previously).
      • If nothing picked → treat the first sheet as combined.
      • The combined extractor contains an Unibet-specific branch that only triggers
        when Unibet headers are present; otherwise it falls back to a generic header-scan
        (PokerStars/Bwin friendly).
    """
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = excel_file.wb
    rows = []

    sel_dep = set(getattr(excel_file, "selected_deposit_sheets", []) or [])
    sel_wd = set(getattr(excel_file, "selected_cashout_sheets", []) or [])

    # If nothing selected → treat first sheet as combined
    if not sel_dep and not sel_wd:
        sheet = wb[wb.sheetnames[0]]
        rows += extract_rows_from_combined_sheet(sheet, excel_file, log)

    else:
        both = sel_dep & sel_wd
        if both:
            log(
                f"ℹ️ Same sheet in both lists → will parse once as combined: {sorted(both)}"
            )
            for name in sorted(both):
                rows += extract_rows_from_combined_sheet(wb[name], excel_file, log)

        # IMPORTANT: for distinct deposit/cashout sheets, use the old per-sheet path
        for name in sorted(sel_dep - both):
            rows += [
                # convert per-sheet rows [date, type, amount(+/-), status]
                # to the canonical combined layout [Currency, Date, Type, Amount, Status]
                [None, r[0], "Deposit", float(abs(r[2])), r[3]]
                for r in extract_rows_from_sheet(
                    wb[name], "Deposit", "dep_settled", log
                )
            ]

        for name in sorted(sel_wd - both):
            rows += [
                [None, r[0], "Redeem", float(abs(r[2])), r[3]]
                for r in extract_rows_from_sheet(wb[name], "Redeem", "pay_closed", log)
            ]

    # (Re)create Payment Summary
    if "Payment Summary" in wb.sheetnames:
        del wb["Payment Summary"]
    ws = wb.create_sheet("Payment Summary")

    headers = ["Currency", "Date", "Type", "Amount", "Status"]
    ws.append(headers)
    for r in rows:
        ws.append(r)

    if rows:
        end_col = get_column_letter(len(headers))
        ref = f"A1:{end_col}{len(rows) + 1}"
        tbl = Table(displayName="PaymentSummary", ref=ref)
        tbl.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium9", showRowStripes=True, showColumnStripes=False
        )
        ws.add_table(tbl)

    # Basic widths
    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 16

    dep_sum = sum(r[3] for r in rows if r[2] == "Deposit")
    wd_sum = sum(r[3] for r in rows if r[2] == "Redeem")
    net = dep_sum - wd_sum
    log(f"TA logic: Deposits = {dep_sum}, Cashouts = {wd_sum}, Net = {net}")
    log(f"✅ Created Payment Summary with {len(rows)} rows.")


def _ecb_average_rate(base: str, quote: str, start_date, end_date, log) -> float:
    """
    Fetch average daily ECB reference rate for [start_date, end_date] inclusive.

    ECB publishes series as: EXR/D.<CURRENCY>.EUR.SP00.A  (<currency> per 1 EUR).
    To convert BASE -> QUOTE:
      - If BASE == QUOTE: 1.0
      - If BASE == 'EUR' and QUOTE != 'EUR': use QUOTE/EUR directly.
      - If QUOTE == 'EUR' and BASE != 'EUR': use BASE/EUR and INVERT.
      - Else: cross via EUR: (QUOTE/EUR) * (1 / (BASE/EUR))

    Nearest fallback:
      - If the requested span has no data (e.g. weekend/holiday), search the nearest
        available single day (for single-day spans) or expand the window symmetrically
        (for multi-day spans) up to ±nearest_days until data is found.

    Returns the multiplier to convert amounts in BASE to QUOTE (float).
    """
    import requests
    from datetime import datetime, timedelta
    import time

    nearest_days = 7  # how far to search when the requested range has no data

    def _d(d):
        if isinstance(d, datetime):
            return d.strftime("%Y-%m-%d")
        return str(d)[:10]

    s_str, e_str = _d(start_date), _d(end_date)
    base = (base or "").upper()
    quote = (quote or "").upper()

    if base == quote:
        log(f"🌍 FX {base}→{quote} {s_str}…{e_str}: identical currencies, avg=1.000000")
        return 1.0

    cache_key = (base, quote, s_str, e_str)
    if not hasattr(_ecb_average_rate, "_cache"):
        _ecb_average_rate._cache = {}
    if cache_key in _ecb_average_rate._cache:
        return _ecb_average_rate._cache[cache_key]

    # --- helpers to fetch <CUR>/EUR series and average it ---
    def _parse_sdmx_json(js):
        data = js.get("dataSets") or js.get("DataSets")
        if not data:
            data = (js.get("data") or {}).get("dataSets")
        if not data:
            return []
        ds0 = data[0]
        series_all = ds0.get("series") or ds0.get("Series")
        vals = []
        if series_all:
            for _, sdict in series_all.items():
                obs = sdict.get("observations") or sdict.get("Observations") or {}
                for _, v in obs.items():
                    raw = v[0] if isinstance(v, list) and v else v
                    try:
                        vals.append(float(raw))
                    except Exception:
                        continue
        else:
            obs = ds0.get("observations") or ds0.get("Observations") or {}
            for _, v in obs.items():
                raw = v[0] if isinstance(v, list) and v else v
                try:
                    vals.append(float(raw))
                except Exception:
                    continue
        return vals

    def _parse_csv(txt: str):
        import csv as _csv
        from io import StringIO

        vals = []
        f = StringIO(txt)
        reader = _csv.DictReader(f)
        for row in reader:
            if not row:
                continue
            v = row.get("OBS_VALUE") or row.get("ObsValue") or row.get("obs_value")
            if v is None or str(v).strip() == "":
                continue
            try:
                vals.append(float(str(v).strip()))
            except Exception:
                continue
        return vals

    def _do_request(url, params, headers, attempts=2, sleep_sec=0.6):
        last_exc = None
        for _ in range(attempts):
            try:
                resp = requests.get(url, params=params, headers=headers, timeout=20)
                return resp
            except Exception as ex:
                last_exc = ex
                time.sleep(sleep_sec)
        if last_exc:
            raise last_exc

    def _avg_series_cur_per_eur(cur_code: str):
        """
        Return (avg, mode, n_points) for EXR/D.<cur_code>.EUR.SP00.A over [start_date,end_date].
        If no data points for the requested span, search nearest within ±nearest_days.
        """
        cur_code = (cur_code or "").upper()

        def __try_span(s, e, announce_log=True):
            s_local, e_local = _d(s), _d(e)
            tries = [
                (
                    "json",
                    f"https://data-api.ecb.europa.eu/service/data/EXR/D.{cur_code}.EUR.SP00.A",
                    {
                        "startPeriod": s_local,
                        "endPeriod": e_local,
                        "format": "sdmx-json",
                    },
                    {
                        "Accept": "application/vnd.sdmx.data+json;version=1.0.0",
                        "User-Agent": "ExcelProcessor/1.0",
                    },
                ),
                (
                    "csv",
                    f"https://data-api.ecb.europa.eu/service/data/EXR/D.{cur_code}.EUR.SP00.A",
                    {"startPeriod": s_local, "endPeriod": e_local, "format": "csvdata"},
                    {
                        "Accept": "text/csv; charset=UTF-8",
                        "User-Agent": "ExcelProcessor/1.0",
                    },
                ),
                (
                    "json",
                    f"https://sdw-wsrest.ecb.europa.eu/service/data/EXR/D.{cur_code}.EUR.SP00.A",
                    {
                        "startPeriod": s_local,
                        "endPeriod": e_local,
                        "format": "sdmx-json",
                    },
                    {
                        "Accept": "application/vnd.sdmx.data+json;version=1.0.0",
                        "User-Agent": "ExcelProcessor/1.0",
                    },
                ),
                (
                    "csv",
                    f"https://sdw-wsrest.ecb.europa.eu/service/data/EXR/D.{cur_code}.EUR.SP00.A",
                    {"startPeriod": s_local, "endPeriod": e_local, "format": "csvdata"},
                    {
                        "Accept": "text/csv; charset=UTF-8",
                        "User-Agent": "ExcelProcessor/1.0",
                    },
                ),
            ]
            last_err = None
            for mode, url, params, headers in tries:
                try:
                    resp = _do_request(url, params, headers)
                    ctype = (resp.headers.get("Content-Type") or "").lower()
                    if resp.status_code != 200:
                        snippet = (resp.text or "")[:120].replace("\n", " ")
                        log(
                            f"🔁 ECB {mode.upper()} {resp.status_code} {ctype} — {snippet}"
                        )
                        continue
                    if mode == "json":
                        try:
                            js = resp.json()
                        except Exception as jerr:
                            snippet = (resp.text or "")[:120].replace("\n", " ")
                            log(
                                f"🔁 ECB JSON parse failed ({ctype}) — {jerr}; body: {snippet}"
                            )
                            continue
                        vals = _parse_sdmx_json(js)
                    else:
                        txt = resp.text or ""
                        vals = _parse_csv(txt)
                    if not vals:
                        snippet = (resp.text or "")[:120].replace("\n", " ")
                        log(
                            f"🔁 ECB parsed 0 points from {mode.upper()} — {ctype}; body: {snippet}"
                        )
                        continue
                    avg = sum(vals) / len(vals)
                    if announce_log:
                        log(
                            f"🌍 FX {cur_code}/EUR {s_local}…{e_local} via {mode.upper()}: {len(vals)} pts, avg={avg:.6f}"
                        )
                    return avg, mode, len(vals)
                except Exception as fetch_err:
                    last_err = fetch_err
                    log(f"🔁 ECB {mode.upper()} attempt failed — {fetch_err}")
            raise RuntimeError(
                f"No data for {cur_code}/EUR {s_local}…{e_local}: {last_err or 'no data'}"
            )

        # 1) Try the requested span
        try:
            return __try_span(start_date, end_date, announce_log=True)
        except Exception as initial_err:
            # 2) Nearest-day / expanded-window fallback
            log(
                f"⚠️ No data for {cur_code}/EUR {s_str}…{e_str}. Searching nearest within ±{nearest_days}d…"
            )
            single_day = s_str == e_str

            if single_day:
                # Search nearest single day outward
                for delta in range(1, nearest_days + 1):
                    for cand in (
                        start_date - timedelta(days=delta),
                        start_date + timedelta(days=delta),
                    ):
                        try:
                            avg, mode, n = __try_span(cand, cand, announce_log=False)
                            log(
                                f"🌍 FX {cur_code}/EUR using nearest {_d(cand)} (Δ{delta}d, {mode.upper()}): {n} pts, avg={avg:.6f}"
                            )
                            return avg, mode, n
                        except Exception:
                            continue
            else:
                # Expand the window symmetrically
                for delta in range(1, nearest_days + 1):
                    s2 = start_date - timedelta(days=delta)
                    e2 = end_date + timedelta(days=delta)
                    try:
                        avg, mode, n = __try_span(s2, e2, announce_log=False)
                        log(
                            f"🌍 FX {cur_code}/EUR expanded to {_d(s2)}…{_d(e2)} (±{delta}d, {mode.upper()}): {n} pts, avg={avg:.6f}"
                        )
                        return avg, mode, n
                    except Exception:
                        continue

            # 3) Give up with the original error context
            raise RuntimeError(
                f"No data for {cur_code}/EUR {s_str}…{e_str}: {initial_err}"
            )

    # --- compute BASE->QUOTE using EUR as necessary ---
    try:
        if base == "EUR":
            avg_quote_per_eur, mode_q, n_q = _avg_series_cur_per_eur(quote)
            rate = avg_quote_per_eur  # EUR -> QUOTE
            log(
                f"🌍 FX {base}→{quote} {s_str}…{e_str}: derived from {quote}/EUR "
                f"(mode {mode_q.upper()}, {n_q} pts), avg={rate:.6f}"
            )
        elif quote == "EUR":
            avg_base_per_eur, mode_b, n_b = _avg_series_cur_per_eur(base)
            rate = 1.0 / avg_base_per_eur  # BASE -> EUR
            log(
                f"🌍 FX {base}→{quote} {s_str}…{e_str}: 1/( {base}/EUR ) "
                f"(mode {mode_b.upper()}, {n_b} pts) => avg={rate:.6f}"
            )
        else:
            avg_base_per_eur, mode_b, n_b = _avg_series_cur_per_eur(base)
            avg_quote_per_eur, mode_q, n_q = _avg_series_cur_per_eur(quote)
            rate = (avg_quote_per_eur) * (1.0 / avg_base_per_eur)
            log(
                f"🌍 FX {base}→{quote} {s_str}…{e_str}: ( {quote}/EUR ) * ( 1/( {base}/EUR ) ) "
                f"(modes {mode_q.upper()}/{mode_b.upper()}, pts {n_q}/{n_b}) => avg={rate:.6f}"
            )
    except Exception as err:
        raise RuntimeError(
            f"ECB fetch failed for {base}->{quote} {s_str}…{e_str}: {err}"
        )

    _ecb_average_rate._cache[cache_key] = rate
    return rate


def create_payment_summary_eur(excel_file, log):
    """
    Create 'Payment Summary EUR' by converting USD rows in 'Payment Summary' to EUR.
    - Finds USD transactions per YEAR, determines [min(date), max(date)] for that year,
      fetches average USD→EUR from ECB over that span, and applies it.
    - If 'Currency' column is missing, does nothing (by request).
    - EUR rows are copied unchanged. Other currencies are copied unchanged (not converted).
    """
    from datetime import datetime
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = excel_file.wb
    if "Payment Summary" not in wb.sheetnames:
        log("ℹ️ No 'Payment Summary' sheet found — skipping EUR conversion.")
        return

    src = wb["Payment Summary"]
    headers = [str(c.value or "").strip() for c in src[1]]
    try:
        cur_idx = headers.index("Currency")
    except ValueError:
        # Spec: if original table has no currency column, skip entirely.
        log("ℹ️ 'Payment Summary' has no 'Currency' column — FX conversion skipped.")
        return

    try:
        date_idx = headers.index("Date")
        type_idx = headers.index("Type")
        amt_idx = headers.index("Amount")
        status_idx = headers.index("Status")
    except ValueError:
        log("⚠️ Unexpected 'Payment Summary' headers — FX conversion skipped.")
        return

    # Gather USD periods per year
    usd_year_bounds = {}  # year -> [min_dt, max_dt]
    rows = list(src.iter_rows(min_row=2, values_only=True))
    for r in rows:
        cur = (r[cur_idx] or "").strip().upper()
        dt = r[date_idx]
        if not isinstance(dt, datetime):
            continue
        if cur == "USD":
            y = dt.year
            if y not in usd_year_bounds:
                usd_year_bounds[y] = [dt, dt]
            else:
                if dt < usd_year_bounds[y][0]:
                    usd_year_bounds[y][0] = dt
                if dt > usd_year_bounds[y][1]:
                    usd_year_bounds[y][1] = dt

    # Fetch averages per year (once each)
    usd_year_rates = {}
    for y, (d_min, d_max) in sorted(usd_year_bounds.items()):
        try:
            rate = _ecb_average_rate("USD", "EUR", d_min, d_max, log)
            usd_year_rates[y] = rate
            log(
                f"🌍 FX USD→EUR {y} ({d_min:%Y-%m-%d}…{d_max:%Y-%m-%d}) avg = {rate:.6f}"
            )
        except Exception as e:
            log(f"⚠️ ECB fetch failed for {y} ({d_min:%Y-%m-%d}…{d_max:%Y-%m-%d}): {e}")

    # Write new sheet
    if "Payment Summary EUR" in wb.sheetnames:
        del wb["Payment Summary EUR"]
    ws = wb.create_sheet("Payment Summary EUR")

    out_headers = ["Currency", "Date", "Type", "Amount", "Status"]
    ws.append(out_headers)

    converted_rows = 0
    for r in rows:
        cur = (r[cur_idx] or "").strip().upper()
        dt = r[date_idx]
        typ = r[type_idx]
        amt = r[amt_idx]
        st = r[status_idx]

        if isinstance(amt, (int, float)):
            amount = float(amt)
        else:
            # simple guard; amounts in this sheet should already be numeric
            try:
                amount = float(str(amt).replace(",", "."))
            except Exception:
                continue

        if isinstance(dt, datetime) and cur == "USD" and dt.year in usd_year_rates:
            eur_amt = round(amount * float(usd_year_rates[dt.year]), 2)
            ws.append(["EUR", dt, typ, eur_amt, st])
            converted_rows += 1
        else:
            # copy EUR as-is; other currencies passed through unchanged
            if cur == "EUR":
                ws.append(["EUR", dt, typ, amount, st])
            else:
                # Keep original currency/amount (not converted)
                ws.append([cur, dt, typ, amount, st])

    # Table styling
    if ws.max_row > 1:
        end_col = get_column_letter(len(out_headers))
        ref = f"A1:{end_col}{ws.max_row}"
        tbl = Table(displayName="PaymentSummaryEUR", ref=ref)
        tbl.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium9", showRowStripes=True, showColumnStripes=False
        )
        ws.add_table(tbl)

    # Light formatting (numbers/dates) best-effort
    try:
        # Amount col
        amt_col = out_headers.index("Amount") + 1
        for cell in ws.iter_cols(min_col=amt_col, max_col=amt_col, min_row=2)[0]:
            cell.number_format = "#,##0.00"
        # Date col
        dt_col = out_headers.index("Date") + 1
        for cell in ws.iter_cols(min_col=dt_col, max_col=dt_col, min_row=2)[0]:
            cell.number_format = "DD.MM.YYYY"
    except Exception:
        pass

    log(f"✅ Created 'Payment Summary EUR' ({converted_rows} USD rows converted).")


def _create_currency_type_year_summary_from(
    excel_file, src_sheet_name, out_sheet_name, log
):
    """
    Build a compact grouped summary from `src_sheet_name` into `out_sheet_name`,
    aggregating by (Currency, Type, Year) with sum(Amount).
    Only numeric amounts are considered. Dates are expected to be datetime, but
    strings are tolerated best-effort.

    Columns expected in source:
      ["Currency", "Date", "Type", "Amount", "Status"]
    """
    from datetime import datetime
    from collections import defaultdict
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter

    wb = excel_file.wb
    if src_sheet_name not in wb.sheetnames:
        log(f"ℹ️ No '{src_sheet_name}' sheet found — skipping '{out_sheet_name}'.")
        return

    src = wb[src_sheet_name]

    # map header -> index
    header = [str(c.value or "").strip() for c in src[1]]
    try:
        idx_cur = header.index("Currency")
        idx_dt = header.index("Date")
        idx_typ = header.index("Type")
        idx_amt = header.index("Amount")
    except ValueError:
        log(
            f"⚠️ Unexpected headers in '{src_sheet_name}' — skipping '{out_sheet_name}'."
        )
        return

    # aggregate
    agg = defaultdict(float)

    def _coerce_year(v):
        if isinstance(v, datetime):
            return v.year
        # tolerate strings like '2014-11-06 11:35:22'
        s = str(v or "").strip()
        if not s:
            return None
        for fmt in (
            "%Y-%m-%d %H:%M:%S",
            "%Y-%m-%d",
            "%d-%m-%Y",
            "%d/%m/%Y",
            "%Y/%m/%d",
        ):
            try:
                return datetime.strptime(s, fmt).year
            except Exception:
                continue
        return None

    for row in src.iter_rows(min_row=2, values_only=True):
        cur = (row[idx_cur] or "").strip().upper()
        # Only summarize true EUR rows in the EUR sheet
        if cur != "EUR":
            continue

        yr = _coerce_year(row[idx_dt])
        typ = (row[idx_typ] or "").strip()
        amt = row[idx_amt]

        if yr is None or typ == "":
            continue
        try:
            val = float(amt)
        except Exception:
            continue

        agg[(cur, typ, yr)] += val

    # create/replace output sheet
    if out_sheet_name in wb.sheetnames:
        del wb[out_sheet_name]
    ws = wb.create_sheet(out_sheet_name)

    out_headers = ["Currency", "Type", "Year", "Amount"]
    ws.append(out_headers)

    # stable ordering: Currency, Type, Year
    rows_written = 0
    for cur, typ, yr in sorted(agg.keys(), key=lambda k: (k[0], k[1], k[2])):
        ws.append([cur, typ, int(yr), round(agg[(cur, typ, yr)], 2)])
        rows_written += 1

    # style as Excel Table
    if rows_written > 0:
        end_col = get_column_letter(len(out_headers))
        ref = f"A1:{end_col}{ws.max_row}"
        tbl = Table(displayName="SummaryEUR", ref=ref)
        tbl.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium9", showRowStripes=True, showColumnStripes=False
        )
        ws.add_table(tbl)

        # formats
        try:
            # Year column integer
            yr_col = out_headers.index("Year") + 1
            for cell in ws.iter_cols(min_col=yr_col, max_col=yr_col, min_row=2)[0]:
                cell.number_format = "0"
            # Amount column
            amt_col = out_headers.index("Amount") + 1
            for cell in ws.iter_cols(min_col=amt_col, max_col=amt_col, min_row=2)[0]:
                cell.number_format = "#,##0.00"
        except Exception:
            pass

    log(f"✅ Created 'Summary EUR' with {rows_written} rows (Currency → Type → Year).")


def create_summary_sheet_eur(excel_file, log):
    """
    Build a pivot-like 'Summary EUR' from 'Payment Summary EUR':
      Currency
        Deposit
          <Year>
        Withdrawal
          <Year>
      Grand Total

    Assumes source columns: ["Currency", "Date", "Type", "Amount", "Status"].
    'Redeem' is shown as 'Withdrawal' with negative totals.
    """
    from datetime import datetime
    from collections import defaultdict
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter

    wb = excel_file.wb
    src_name = "Payment Summary EUR"
    out_name = "Summary EUR"

    if src_name not in wb.sheetnames:
        log(f"ℹ️ No '{src_name}' sheet found — skipping '{out_name}'.")
        return

    src = wb[src_name]

    # Header mapping
    header = [str(c.value or "").strip() for c in src[1]]
    try:
        idx_cur = header.index("Currency")
        idx_dt = header.index("Date")
        idx_typ = header.index("Type")
        idx_amt = header.index("Amount")
    except ValueError:
        log(f"⚠️ Unexpected headers in '{src_name}' — skipping '{out_name}'.")
        return

    # Helpers
    def _to_dt(x):
        if isinstance(x, datetime):
            return x
        s = str(x or "").strip()
        if not s:
            return None
        for fmt in (
            "%Y-%m-%d %H:%M:%S",
            "%Y-%m-%d",
            "%d-%m-%Y",
            "%d/%m/%Y",
            "%Y/%m/%d",
        ):
            try:
                return datetime.strptime(s, fmt)
            except Exception:
                continue
        return None

    # Aggregate: {currency: {"Deposit": {year: sum}, "Withdrawal": {year: -sum}}}
    agg = defaultdict(
        lambda: {"Deposit": defaultdict(float), "Withdrawal": defaultdict(float)}
    )

    min_dt = None
    max_dt = None

    for row in src.iter_rows(min_row=2, values_only=True):
        cur = (row[idx_cur] or "").strip().upper()
        dt = _to_dt(row[idx_dt])
        typ_raw = (row[idx_typ] or "").strip()
        amt = row[idx_amt]

        if not cur or not isinstance(amt, (int, float)):
            continue
        if dt is None:
            continue

        # date range tracking
        if (min_dt is None) or (dt < min_dt):
            min_dt = dt
        if (max_dt is None) or (dt > max_dt):
            max_dt = dt

        year = dt.year

        # Normalize type to Deposit / Withdrawal
        typ_l = typ_raw.lower()
        if "deposit" in typ_l or typ_l == "credit":
            agg[cur]["Deposit"][year] += float(amt)
        else:
            # treat everything else as Withdrawal (Redeem)
            agg[cur]["Withdrawal"][year] += -abs(float(amt))  # negative

    # Create/replace output
    if out_name in wb.sheetnames:
        del wb[out_name]
    ws = wb.create_sheet(out_name)

    # Title
    title = "Summary of payment transactions"
    if min_dt and max_dt:
        title += f" ({min_dt.strftime('%d.%m.%Y')} - {max_dt.strftime('%d.%m.%Y')})"
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
    ws["A1"].font = Font(bold=True)
    ws.append([])

    # Headers
    ws.append(["Row Labels", "Sum of Amount"])
    ws["A3"].font = Font(bold=True)
    ws["B3"].font = Font(bold=True)

    # Number format for amounts
    amt_fmt = "#,##0.00;[Red]-#,##0.00"

    r = 4
    grand_total = 0.0

    # Stable currency order
    for cur in sorted(agg.keys()):
        dep_years = agg[cur]["Deposit"]
        wdr_years = agg[cur]["Withdrawal"]

        cur_total = sum(dep_years.values()) + sum(wdr_years.values())
        grand_total += cur_total

        # Currency row
        ws.cell(row=r, column=1, value=cur).font = Font(bold=True)
        c_amt = ws.cell(row=r, column=2, value=round(cur_total, 2))
        c_amt.number_format = amt_fmt
        c_amt.font = Font(bold=True)
        r += 1

        # Deposit block
        if dep_years:
            lbl = ws.cell(row=r, column=1, value="Deposit")
            lbl.font = Font(bold=True)
            lbl.alignment = Alignment(indent=1)
            d_sum = round(sum(dep_years.values()), 2)
            d_amt = ws.cell(row=r, column=2, value=d_sum)
            d_amt.number_format = amt_fmt
            d_amt.font = Font(bold=True)
            r += 1

            for yr in sorted(dep_years.keys()):
                y_lbl = ws.cell(row=r, column=1, value=int(yr))
                y_lbl.alignment = Alignment(indent=2)
                y_amt = ws.cell(row=r, column=2, value=round(dep_years[yr], 2))
                y_amt.number_format = amt_fmt
                r += 1

        # Withdrawal block
        if wdr_years:
            lbl = ws.cell(row=r, column=1, value="Withdrawal")
            lbl.font = Font(bold=True)
            lbl.alignment = Alignment(indent=1)
            w_sum = round(sum(wdr_years.values()), 2)  # negative
            w_amt = ws.cell(row=r, column=2, value=w_sum)
            w_amt.number_format = amt_fmt
            w_amt.font = Font(bold=True)
            r += 1

            for yr in sorted(wdr_years.keys()):
                y_lbl = ws.cell(row=r, column=1, value=int(yr))
                y_lbl.alignment = Alignment(indent=2)
                y_amt = ws.cell(row=r, column=2, value=round(wdr_years[yr], 2))
                y_amt.number_format = amt_fmt
                r += 1

        # spacer between currencies
        ws.append([])
        r += 1

    # Grand total row
    ws.cell(row=r, column=1, value="Grand Total").font = Font(bold=True)
    g_amt = ws.cell(row=r, column=2, value=round(grand_total, 2))
    g_amt.font = Font(bold=True)
    g_amt.number_format = amt_fmt

    # Column widths
    ws.column_dimensions[get_column_letter(1)].width = 36
    ws.column_dimensions[get_column_letter(2)].width = 18

    log(f"✅ Created '{out_name}' (pivot-style Currency → Type → Year).")


def extract_rows_from_sheet(ws, tx_label, valid_status, log):
    """
    For separate Deposit/Cashout sheets.
    Returns list of [date, type, amount(+/-), status_str]
    """
    import re, builtins
    from datetime import datetime as _dt

    def _s(x):
        return "" if x is None else builtins.str(x)

    # --- find header row (first non-empty with 4+ unique tokens) ---
    header_row = None
    hmap = {}
    max_scan = min(ws.max_row, 500)
    for r in range(1, max_scan + 1):
        vals = [
            _s(ws.cell(row=r, column=c).value).strip()
            for c in range(1, ws.max_column + 1)
        ]
        norm = [re.sub(r"[^a-z0-9]+", "", v.lower()) for v in vals]
        if len({t for t in norm if t}) >= 4 and any(
            k in norm
            for k in (
                "created",
                "date",
                "datum",
                "timestamp",
                "requested",
                "status",
                "state",
                "credited",
                "txnamount",
                "amount",
                "initial",
            )
        ):
            header_row = r
            for i, h in enumerate(norm):
                if h and h not in hmap:
                    hmap[h] = i
            break

    if not header_row:
        log(f"⚠️ No distinct header row found (scanned {max_scan} rows).")
        return []

    # --- column indices ---
    date_idx = next(
        (
            hmap[k]
            for k in ("created", "date", "datum", "timestamp", "requested", "completed")
            if k in hmap
        ),
        None,
    )
    time_idx = next((hmap[k] for k in ("time", "tijd") if k in hmap), None)
    status_idx = next(
        (
            hmap[k]
            for k in ("status", "state", "ergebnis", "resultado", "settled")
            if k in hmap
        ),
        None,
    )

    amount_idx = None
    for key in ("credited", "txnamount", "amount", "initial", "net", "value"):
        if key in hmap:
            amount_idx = hmap[key]
            break
    if amount_idx is None:
        log(
            "❌ No recognizable amount column (looking for Credited/Txn Amount/Amount/Initial/Net)."
        )
        return []

    # status acceptance (include if status missing or looks 'good')
    good_status_tokens = {
        re.sub(r"[^a-z0-9]+", "", valid_status.lower()) if valid_status else "",
        "dep_settled",
        "pay_closed",
        "settled",
        "completed",
        "complete",
        "success",
        "succeeded",
        "approved",
        "processed",
        "ok",
        "closed",
        "paid",
    }
    bad_status_tokens = {
        "rejected",
        "failed",
        "declined",
        "cancelled",
        "canceled",
        "chargeback",
        "pending",
        "hold",
        "onhold",
    }

    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]

        # status gate (permissive but excludes obvious bads)
        stxt_raw = (
            _s(vals[status_idx]).strip().lower().replace(" ", "")
            if status_idx is not None
            else ""
        )
        if stxt_raw and any(b in stxt_raw for b in bad_status_tokens):
            continue
        if stxt_raw and not any(g in stxt_raw for g in good_status_tokens):
            # not strictly required — still allow, but prefer good statuses
            pass

        # date/time
        dval = vals[date_idx] if date_idx is not None else None
        if isinstance(dval, _dt):
            d_out = dval
        else:
            dtxt = _s(dval)
            ttxt = _s(vals[time_idx]) if time_idx is not None else ""
            combo = (dtxt + " " + ttxt).strip()
            d_out = dval
            for fmt in (
                "%d.%m.%Y %H:%M:%S",
                "%Y-%m-%d %H:%M:%S",
                "%d/%m/%Y %H:%M",
                "%Y/%m/%d %H:%M:%S",
                "%d-%m-%Y %H:%M",
                "%Y-%m-%d",
                "%d/%m/%Y",
                "%d-%m-%Y",
            ):
                try:
                    d_out = _dt.strptime(combo or dtxt, fmt)
                    break
                except Exception:
                    pass

        # amount
        aval = vals[amount_idx]
        if isinstance(aval, (int, float)):
            amount = float(aval)
        else:
            txt = (
                _s(aval)
                .replace("€", "")
                .replace("$", "")
                .replace("£", "")
                .replace(" ", "")
            )
            if "," in txt and "." in txt:
                txt = txt.replace(".", "").replace(",", ".")
            else:
                txt = txt.replace(",", ".")
            try:
                amount = float(txt)
            except Exception:
                continue

        signed = amount if tx_label == "Deposit" else -abs(amount)
        rows.append([d_out, tx_label, signed, stxt_raw])

    return rows


def _is_unibet_sheet(sheet):
    """
    Recognize Unibet by its very specific headers or sheet title.
    Returns True if we should use the Unibet-specific parsing.
    """
    title_hit = (
        str(getattr(sheet, "title", "")).strip().lower() == "transaction history"
    )

    # Look at the first row only; Unibet exports have stable human headers.
    raw_headers = [cell.value for cell in sheet[1]]

    def _h(name):
        return any(
            isinstance(h, str) and h.strip().lower() == name for h in raw_headers
        )

    header_hit = (
        _h("ref date") and _h("amount") and (_h("transaction type name") or _h("type"))
    )
    return title_hit or header_hit


def _parse_chunk_worker(args):
    """
    args = (
        chunk,                # list[(row_idx:int, row:tuple)]
        idx_amount, idx_curr, idx_status, idx_type, idx_ref_date, idx_conf_dt,
        force_type,           # "Deposit"|"Redeem"|None
        date_filter_dict,     # {"mode": "...", "from": dt|None, "to": dt|None} or None
    )
    Returns: list[[Currency, Date(datetime), "Deposit"/"Redeem", Amount(+), Status]]
    """
    from datetime import datetime
    from processor.utils import parse_any_datetime
    import re
    from collections import Counter

    (
        chunk,
        idx_amount,
        idx_curr,
        idx_status,
        idx_type,
        idx_ref_date,
        idx_conf_dt,
        force_type,
        df,
    ) = args

    # ---- helpers (unchanged logic) -----------------------------------------
    def _row_passes_date_filter(dt: datetime) -> bool:
        if not df or not isinstance(dt, datetime):
            return True
        mode = df.get("mode")
        d_from = df.get("from")
        d_to = df.get("to")
        if mode == "before" and d_to:
            return dt.date() <= d_to.date()
        if mode == "after" and d_from:
            return dt.date() >= d_from.date()
        if mode == "between":
            if d_from and d_to and d_from > d_to:
                d_from, d_to = d_to, d_from
            return (not d_from or dt.date() >= d_from.date()) and (
                not d_to or dt.date() <= d_to.date()
            )
        return True

    def _tx_label(raw):
        s = (raw or "").strip().lower()
        s_compact = re.sub(r"[^a-z0-9]+", "", s)

        # --- Redeem / withdrawal patterns (very liberal) ---
        if (
            "withdraw" in s_compact
            or "withdrawal" in s_compact
            or "cashout" in s_compact
            or "cash out" in s
            or "payout" in s_compact
            or "pay out" in s
            or "redeem" in s_compact
            or "debit" in s_compact
            or "transferout" in s_compact
            or "bankwithdrawal" in s_compact
            or s in {"red", "wd"}  # PokerStars shorthands
            or s_compact.startswith(("wd", "withdr"))  # prefixes
        ):
            return "Redeem"

        # --- Deposit patterns ---
        if (
            "deposit" in s_compact
            or "topup" in s_compact
            or "top-up" in s
            or "credit" in s_compact
            or "buyin" in s_compact
            or "buy-in" in s
            or "payment" in s_compact
            or s == "dep"  # PokerStars shorthand
        ):
            return "Deposit"

        if s in {"credit", "credited"}:
            return "Deposit"
        if s in {"debit", "debited"}:
            return "Redeem"
        return None

    def _status_norm(x, label_hint=None):
        s = (str(x) if x is not None else "").strip().lower()
        if not s:
            return "auto"  # accept when status column is missing
        if s in {"dep_settled", "pay_closed", "rdm_closed"}:
            return s
        if any(
            tok in s
            for tok in (
                "approved",
                "completed",
                "settled",
                "closed",
                "success",
                "paid",
                "posted",
                "processed",
                "ok",
                "done",
            )
        ):
            return s
        return s

    # ---- diagnostics -------------------------------------------------------
    # Per-chunk counters (printed at the end of the worker)
    seen_types = Counter()  # raw Type strings
    labeled_counts = Counter()  # "Deposit"/"Redeem"/None before defaulting
    accepted_counts = Counter()  # kept rows per label
    rejected_by_status = Counter()  # rows dropped due to status gate
    skip_reasons = Counter()  # why a row was skipped
    statuses_seen = Counter()  # normalized statuses
    currencies_seen = Counter()

    # upfront config snapshot
    print(
        f"🧩 Worker start: rows={len(chunk)} | "
        f"cols(amount={idx_amount}, curr={idx_curr}, status={idx_status}, "
        f"type={idx_type}, ref_date={idx_ref_date}, conf_date={idx_conf_dt}) | "
        f"force_type={force_type} | date_filter={df}"
    )

    out = []
    for _r_i, row in chunk:
        try:
            # Date (prefer confirmed/completed)
            d1 = row[idx_ref_date] if idx_ref_date is not None else None
            d2 = row[idx_conf_dt] if idx_conf_dt is not None else None
            dt = None
            for cand in (d2, d1):
                if cand in (None, ""):
                    continue
                if isinstance(cand, datetime):
                    dt = cand
                    break
                dt = parse_any_datetime(str(cand))
                if dt:
                    break
            if not isinstance(dt, datetime):
                skip_reasons["no_or_unparsable_date"] += 1
                continue
            if not _row_passes_date_filter(dt):
                skip_reasons["filtered_by_date"] += 1
                continue

            # Amount → positive
            raw_amount = row[idx_amount]
            if raw_amount is None or (
                isinstance(raw_amount, str) and not str(raw_amount).strip()
            ):
                skip_reasons["no_amount"] += 1
                continue
            a_str = str(raw_amount).replace("€", "").replace("\u00a0", " ").strip()
            if "," in a_str and a_str.rfind(",") > a_str.rfind("."):
                a_str = a_str.replace(".", "").replace(",", ".")
            else:
                a_str = a_str.replace(",", "")
            try:
                amount_val = abs(float(a_str))
            except Exception:
                skip_reasons["amount_parse_error"] += 1
                continue

            # Type
            type_raw = row[idx_type] if idx_type is not None else ""
            seen_types[(str(type_raw or "")).strip().lower() or "(blank)"] += 1
            if force_type in {"Deposit", "Redeem"}:
                label = force_type
                labeled_counts[label] += 1
            else:
                inferred = _tx_label(type_raw)
                labeled_counts[inferred or "None"] += 1
                label = inferred or "Deposit"  # default behaviour (unchanged)
                if inferred is None:
                    skip_reasons["defaulted_to_deposit"] += 1

            # Currency / Status
            currency = row[idx_curr] if idx_curr is not None else None
            currencies_seen[str(currency or "").upper() or "—"] += 1
            status_raw = row[idx_status] if idx_status is not None else ""
            norm_status = _status_norm(status_raw, label_hint=label)
            statuses_seen[norm_status] += 1

            # --- force direction by status, if status is decisive ---
            from processor.normalization import normalize_status

            ns_bucket = normalize_status(status_raw)
            if ns_bucket == "accepted_withdraw":
                label = "Redeem"
            elif ns_bucket == "accepted_deposit":
                label = "Deposit"

            s_raw = (str(status_raw) or "").lower()
            if any(
                tok in s_raw
                for tok in ("rdm_closed", "pay_closed", "payout_closed", "payout_paid")
            ):
                label = "Redeem"

            # Extra safety: check normalized status for withdrawal indicators
            ns = normalize_status(status_raw)
            if ns in ("rdm_closed", "pay_closed", "payout_closed", "payout_paid"):
                label = "Redeem"

            # Include rows (accept 'auto' when status missing)
            if label == "Deposit":
                allowed = {
                    "dep_settled",
                    "approved",
                    "completed",
                    "closed",
                    "success",
                    "paid",
                    "posted",
                    "processed",
                    "ok",
                    "done",
                    "auto",
                }
                if norm_status in allowed:
                    out.append(
                        [
                            str(currency or "").upper() or "EUR",
                            dt,
                            "Deposit",
                            float(amount_val),
                            norm_status,
                        ]
                    )
                    accepted_counts["Deposit"] += 1
                else:
                    rejected_by_status["Deposit"] += 1

            else:  # Redeem
                allowed = {
                    "pay_closed",
                    "completed",
                    "closed",
                    "success",
                    "paid",
                    "processed",
                    "posted",
                    "ok",
                    "done",
                    "settled",
                    "auto",
                    "approved",
                }
                if norm_status in allowed:
                    out.append(
                        [
                            str(currency or "").upper() or "EUR",
                            dt,
                            "Redeem",
                            float(amount_val),
                            norm_status,
                        ]
                    )
                    accepted_counts["Redeem"] += 1
                else:
                    rejected_by_status["Redeem"] += 1

        except Exception as e:
            skip_reasons["exception"] += 1
            # include a tiny sample of exceptions to avoid log spam
            if skip_reasons["exception"] <= 5:
                print(f"⚠️ Worker row {_r_i}: exception {e!r}")

    # ---- end-of-chunk diagnostics ------------------------------------------
    # Top samples to keep logs readable
    def _fmt_top(counter: Counter, n=10):
        return ", ".join(f"{k}×{v}" for k, v in counter.most_common(n)) or "—"

    print("🔎 Worker sample types (top 10): " + _fmt_top(seen_types, 10))
    print(
        f"🧮 Worker labeling: "
        f"Deposit labeled={labeled_counts.get('Deposit', 0)}, "
        f"Redeem labeled={labeled_counts.get('Redeem', 0)}, "
        f"defaulted_to_deposit={skip_reasons.get('defaulted_to_deposit', 0)}"
    )
    print(
        f"✅ Accepted → Deposit={accepted_counts.get('Deposit', 0)}, "
        f"Redeem={accepted_counts.get('Redeem', 0)}; "
        f"🚫 Rejected by status → "
        f"Deposit={rejected_by_status.get('Deposit', 0)}, "
        f"Redeem={rejected_by_status.get('Redeem', 0)}"
    )
    print(
        "⛔ Skips: "
        + ", ".join(
            f"{k}={v}"
            for k, v in (
                ("no_or_unparsable_date", skip_reasons.get("no_or_unparsable_date", 0)),
                ("filtered_by_date", skip_reasons.get("filtered_by_date", 0)),
                ("no_amount", skip_reasons.get("no_amount", 0)),
                ("amount_parse_error", skip_reasons.get("amount_parse_error", 0)),
                ("exception", skip_reasons.get("exception", 0)),
            )
        )
    )
    print("💱 Currencies seen: " + _fmt_top(currencies_seen, 10))
    print("📌 Statuses (top 10): " + _fmt_top(statuses_seen, 10))
    print(f"🧩 Worker end: produced {len(out)} rows\n")

    return out


def extract_rows_from_combined_sheet(sheet, excel_file, log, force_type=None):
    """
    Universal extractor for a single sheet.
    Chunked single-process parsing (no external worker dependency).

    Returns rows as:
      [Currency(str), Date(datetime), "Deposit"/"Redeem", Amount(float +), Status(str)]

    Filtering rules:
      - If a Status column exists:
          * When values use A/C/P/W: KEEP ONLY A (Approved). Drop C/P/W.
          * Otherwise, use broad final/non-final detection (unchanged).
      - If NO Status column exists:
          * Accept rows with status 'auto'.
    """
    import re
    from datetime import datetime
    from processor.utils import normalize_str, parse_any_datetime

    _TRACE = bool(getattr(excel_file, "trace_rows", True))

    def t(msg):
        if _TRACE:
            try:
                log(f"🔎 {msg}")
            except Exception:
                pass

    # ---------------- helpers ----------------
    def _row_passes_date_filter(dt):
        df = getattr(excel_file, "date_filter", None)
        if not df or not isinstance(dt, datetime):
            return True
        mode = df.get("mode")
        d_from = df.get("from")
        d_to = df.get("to")
        if mode == "before" and d_to:
            return dt.date() <= d_to.date()
        if mode == "after" and d_from:
            return dt.date() >= d_from.date()
        if mode == "between":
            if d_from and d_to and d_from > d_to:
                d_from, d_to = d_to, d_from
            return (not d_from or dt.date() >= d_from.date()) and (
                not d_to or dt.date() <= d_to.date()
            )
        return True

    def _norm_header_token(s):
        if s is None:
            return ""
        return re.sub(r"[^a-z0-9]+", "", str(s).strip().lower())

    # --- robust date picker (normalization.parse_date + Excel serials + fallback) ---
    def _pick_date(d_candidates):
        """
        Return the first valid datetime from a list of candidate cells.
        Supports:
          - native datetime objects
          - Excel serial numbers (days since 1899-12-30) — but ONLY if in a realistic range
          - explicit common US/EU string formats
          - normalization.parse_date
          - utils.parse_any_datetime
        """
        from datetime import datetime, timedelta

        try:
            from processor.normalization import parse_date as _norm_parse_date
        except Exception:
            _norm_parse_date = None
        try:
            from processor.utils import parse_any_datetime as _utils_parse_any_dt
        except Exception:
            _utils_parse_any_dt = None

        # Common explicit formats we want to recognize eagerly
        _MANUAL_FMTS = (
            # date only
            "%m/%d/%Y",
            "%d/%m/%Y",
            "%Y-%m-%d",
            "%d-%m-%Y",
            "%d.%m.%Y",
            # date + time (24h)
            "%m/%d/%Y %H:%M",
            "%m/%d/%Y %H:%M:%S",
            "%d/%m/%Y %H:%M",
            "%d/%m/%Y %H:%M:%S",
            "%Y-%m-%d %H:%M",
            "%Y-%m-%d %H:%M:%S",
            # date + time (AM/PM)
            "%m/%d/%Y %I:%M %p",
            "%m/%d/%Y %I:%M:%S %p",
            "%d/%m/%Y %I:%M %p",
            "%d/%m/%Y %I:%M:%S %p",
        )

        for candidate in d_candidates:
            if candidate in (None, ""):
                continue

            # already a datetime
            if isinstance(candidate, datetime):
                return candidate

            # Excel serial date/time — ONLY for realistic ranges
            # Excel's "zero" is 1899-12-30; serial ~20000 starts in 1954; ~60000 is 2064.
            if isinstance(candidate, (int, float)):
                try:
                    candidate_val = float(candidate)
                    if 20000 <= candidate_val <= 60000:
                        base = datetime(1899, 12, 30)
                        result_date = base + timedelta(days=int(candidate_val))
                        return result_date
                except Exception:
                    pass

            # Treat as string
            s = str(candidate).strip()
            if not s:
                continue

            # Try explicit formats first (fast, predictable)
            for fmt in _MANUAL_FMTS:
                try:
                    return datetime.strptime(s, fmt)
                except Exception:
                    pass

            # normalization.parse_date (your DATE_FORMATS incl. AM/PM)
            if _norm_parse_date:
                try:
                    dt = _norm_parse_date(s)
                    if dt:
                        return dt
                except Exception:
                    pass

            # fallback parser
            if _utils_parse_any_dt:
                try:
                    dt2 = _utils_parse_any_dt(s)
                    if dt2:
                        return dt2
                except Exception:
                    pass

        return None

    def _amount_parse(raw):
        if raw is None:
            return None, 0
        if isinstance(raw, (int, float)):
            val = float(raw)
            return abs(val), -1 if val < 0 else (1 if val > 0 else 0)

        s = str(raw)
        if not s.strip():
            return None, 0

        s = s.replace("\u00a0", " ").strip()

        neg = False
        if s.startswith("(") and s.endswith(")"):
            neg = True
            s = s[1:-1].strip()
        if s.endswith("-"):
            neg = True
            s = s[:-1].strip()
        if s.startswith("-"):
            neg = True
            s = s[1:].strip()

        s = re.sub(r"[A-Za-z$€£¥₽₺₹₩₪₿₫฿₴₦₲៛₡₱ƒ₵₸₭₮₼₨₥₤₢₣₠₧₯₳₰]", "", s)
        s = s.replace(" ", "").replace("'", "")

        if "," in s and "." in s:
            if s.rfind(",") > s.rfind("."):
                s = s.replace(".", "")
                s = s.replace(",", ".")
            else:
                s = s.replace(",", "")
        elif "," in s:
            if s.count(",") == 1:
                s = s.replace(",", ".")
            else:
                s = s.replace(",", "")

        try:
            val = float(s)
        except Exception:
            return None, 0

        if neg:
            val = -val
        return abs(val), -1 if val < 0 else (1 if val > 0 else 0)

    OUT_TOKENS = {
        "withdraw",
        "withdrawal",
        "cashout",
        "cash out",
        "redeem",
        "payout",
        "pay out",
        "refund",
        "chargeback",
        "reversal",
        "reversed",
        "pay_closed",
        "transferout",
        "transfer out",
        "debit",
        "debitmemo",
        "auszahlung",
        "abbuchung",
        "rueckbuchung",
        "rückbuchung",
        "retrait",
        "remboursement",
        "contrepassation",
        "retiro",
        "saque",
        "prelievo",
        "reembolso",
        "rimborso",
        "estorno",
        "estornado",
        "opname",
        "terugbetaling",
        "wypłata",
        "wyplata",
        "obciążenie",
        "вывод",
        "spisanie",
        "cekim",
        "çekim",
    }
    IN_TOKENS = {
        "deposit",
        "topup",
        "top-up",
        "credit",
        "funding",
        "purchase",
        "cardverification",
        "addfunds",
        "add funds",
        "transferin",
        "transfer in",
        "einzahlung",
        "gutschrift",
        "dépôt",
        "depot",
        "versement",
        "depósito",
        "deposito",
        "ingreso",
        "credito",
        "crédito",
        "carte",
        "storting",
        "bijschrijving",
        "wpłata",
        "wplata",
        "ввод",
        "зачисление",
        "yatirma",
        "yatırım",
        "yukleme",
        "yükleme",
    }

    def _tx_label_basic(raw):
        s = raw or ""
        if not s:
            return None
        s_strip = s.strip().lower()
        tkn = _norm_header_token(s_strip)
        if tkn in {"dep", "depo"}:
            return "Deposit"
        if tkn in {"red", "wd", "wdr"}:
            return "Redeem"
        for tok in OUT_TOKENS:
            if tok in tkn:
                return "Redeem"
        for tok in IN_TOKENS:
            if tok in tkn:
                return "Deposit"
        if tkn in {"credited", "credit"}:
            return "Deposit"
        if tkn in {"debited", "debit"}:
            return "Redeem"
        return None

    def _classify_label(
        type_raw, status_raw, gateway_raw, descr_raw, drcr_raw, amount_sign
    ):
        lbl = _tx_label_basic(type_raw)
        if lbl:
            return lbl
        hay = (
            " ".join(
                [
                    str(type_raw or ""),
                    str(status_raw or ""),
                    str(gateway_raw or ""),
                    str(descr_raw or ""),
                    str(drcr_raw or ""),
                ]
            )
            .strip()
            .lower()
        )
        hay_t = _norm_header_token(hay)
        for tok in OUT_TOKENS:
            if tok in hay_t:
                return "Redeem"
        for tok in IN_TOKENS:
            if tok in hay_t:
                return "Deposit"
        dc_t = _norm_header_token(str(drcr_raw or "").strip().lower())
        if dc_t in {"debit", "dr", "d"}:
            return "Redeem"
        if dc_t in {"credit", "cr", "c"}:
            return "Deposit"
        if amount_sign < 0:
            return "Redeem"
        if amount_sign > 0:
            return "Deposit"
        return "Deposit"

    def _status_norm(x):
        s = (str(x) if x is not None else "").strip().lower()
        return s if s else "auto"

    def _is_bad_status(s: str) -> bool:
        if not s:
            return False
        s = s.lower()
        negative = (
            "reject" in s
            or "declin" in s
            or "fail" in s
            or "error" in s
            or "cancel" in s
            or "void" in s
            or "chargeback" in s
            or "charge back" in s
            or "reverse" in s
            or "reversal" in s
            or "refund" in s
            or "refunded" in s
            or "dispute" in s
            or "expired" in s
            or "timeout" in s
            or "time out" in s
            or "refuse" in s
            or "blocked" in s
            or "not processed" in s
            or "insufficient funds" in s
            or "insufficientfunds" in s
            or "denied" in s
        )
        pending = (
            "pending" in s
            or "in review" in s
            or "under review" in s
            or "processing" in s
            or "in process" in s
            or "on hold" in s
            or "await" in s
            or "authorized" in s
            or "authorised" in s
        )
        return negative or pending

    OK_TOKENS_GENERIC = (
        "approved",
        "completed",
        "settled",
        "closed",
        "success",
        "successful",
        "paid",
        "posted",
        "processed",
        "ok",
        "done",
        "confirmed",
        "accepted",
        "captured",
    )
    OK_TOKENS_SPECIAL = {
        "dep_settled",
        "pay_closed",
        "rdm_closed",
        "payout_paid",
        "payout_closed",
    }

    def _is_final_success(s: str) -> bool:
        if not s:
            return False
        s = s.lower()
        if s in OK_TOKENS_SPECIAL:
            return True
        return any(tok in s for tok in OK_TOKENS_GENERIC)

    # compact A/C/P/W mapping
    STATUS_ACPW = {"A": "approved", "C": "canceled", "P": "pending", "W": "waiting"}

    # section divider heuristic
    def _is_section_divider(row_vals):
        if not row_vals:
            return None
        first = row_vals[0]
        if not isinstance(first, str):
            return None
        if not first.strip():
            return None
        # rest mostly empty?
        non_empty_tail = sum(
            1
            for v in row_vals[1:]
            if (isinstance(v, str) and v.strip()) or (v not in (None, ""))
        )
        if non_empty_tail > 0:
            return None
        txt = normalize_str(first)
        # avoid header-like tokens
        if any(
            k in txt
            for k in (
                "date",
                "time",
                "amount",
                "currency",
                "type",
                "transid",
                "gateway",
                "descriptor",
            )
        ):
            return None
        return first.strip()

    # -------------- header detection --------------
    hdr_row, hdr_vals = _find_header_row(sheet, log)
    if hdr_row:
        raw_headers = hdr_vals
        start_data_row = hdr_row + 1
    else:
        raw_headers = [cell.value for cell in sheet[1]]
        start_data_row = 2

    tokens = [_norm_header_token(h) for h in raw_headers]
    t(f"Header row {hdr_row or 1}: {raw_headers}")

    def _find(cands):
        for i, tok in enumerate(tokens):
            if tok in cands:
                return i
        return None

    AMOUNT_CANDS = {
        "amount",
        "amt",
        "value",
        "net",
        "grossamount",
        "netamount",
        "billingamount",
        "txnamount",
        "amountlocal",
        "amountbase",
        "totalamount",
        "transactionamount",
        "amountusd",
        "amounteur",
        "amountgbp",
    }
    CURRENCY_CANDS = {
        "currency",
        "curr",
        "txncurrency",
        "billingcurrency",
        "accountcurrency",
        "currencyid",
    }
    STATUS_CANDS = {"status", "state", "result", "approvalstatus", "paymentstatus"}
    TYPE_CANDS = {
        "type",
        "transtype",
        "transactiontype",
        "transactiontypename",
        "activity",
        "operation",
        "event",
        "code",
        "descriptiontype",
    }
    DATE_REF_CANDS = {
        "date",
        "datetime",
        "timestamp",
        "started",
        "createdon",
        "transdate",
        "valuedate",
        "bookingdate",
        "posteddate",
        "created",
        "startdate",
        "time",
    }
    DATE_CONF_CANDS = {
        "completed",
        "confirmeddate",
        "confirmed",
        "settled",
        "posted",
        "processedon",
        "enddate",
    }
    GATEWAY_CANDS = {
        "gateway",
        "processor",
        "acquirer",
        "provider",
        "method",
        "paymentmethod",
        "channel",
    }
    DESCR_CANDS = {
        "description",
        "details",
        "memo",
        "remark",
        "note",
        "reference",
        "narrative",
        "loggedinreference",
        "loggedinref",
        "loggedin",
        "reason",
    }
    DRCR_CANDS = {"drcr", "debitcredit", "dc", "direction"}

    idx_amount = _find(AMOUNT_CANDS)
    idx_curr = _find(CURRENCY_CANDS)
    idx_status = _find(STATUS_CANDS)
    idx_type = _find(TYPE_CANDS)
    idx_ref_date = _find(DATE_REF_CANDS)
    idx_conf_dt = _find(DATE_CONF_CANDS)
    idx_gateway = _find(GATEWAY_CANDS)
    idx_descr = _find(DESCR_CANDS)
    idx_drcr = _find(DRCR_CANDS)

    log(
        f"🧭 Columns: amount={idx_amount}, currency={idx_curr}, type={idx_type}, "
        f"ref_date={idx_ref_date}, conf_date={idx_conf_dt}, status={idx_status}, "
        f"gateway={idx_gateway}, descr={idx_descr}, drcr={idx_drcr}"
    )

    if idx_amount is None or (idx_ref_date is None and idx_conf_dt is None):
        log("❌ No recognizable amount/date columns (fallbacks failed).")
        return []

    had_status_col = idx_status is not None

    # -------------- iterate rows (chunked) --------------
    rows = []
    CHUNK = 10_000
    max_row = sheet.max_row or 0
    cur = start_data_row

    current_section = None  # track current payment method banner

    while cur <= max_row:
        end = min(max_row, cur + CHUNK - 1)
        for r_i, r in enumerate(
            sheet.iter_rows(min_row=cur, max_row=end, values_only=True), start=cur
        ):
            try:
                # section divider?
                sect = _is_section_divider(r)
                if sect:
                    current_section = sect
                    t(f"row {r_i}: 📑 section = {current_section!r}")
                    continue  # not a data row

                # dates: prefer confirmed/posted over created/started
                d1 = r[idx_ref_date] if idx_ref_date is not None else None
                d2 = r[idx_conf_dt] if idx_conf_dt is not None else None
                dt = _pick_date([d2, d1])
                if not isinstance(dt, datetime):
                    t(f"row {r_i}: ⏭️ skip — bad date ref={d1!r} conf={d2!r}")
                    continue
                if not _row_passes_date_filter(dt):
                    t(f"row {r_i}: ⏭️ skip — date filter failed ({dt.date()})")
                    continue

                raw_amount = r[idx_amount]
                amount_val, sign = _amount_parse(raw_amount)
                if amount_val is None:
                    t(f"row {r_i}: ⏭️ skip — amount parse fail raw={raw_amount!r}")
                    continue

                currency = (r[idx_curr] if idx_curr is not None else None) or ""
                currency = str(currency).upper().strip() or "EUR"
                type_raw = r[idx_type] if idx_type is not None else ""
                status_raw = r[idx_status] if idx_status is not None else ""
                gateway = r[idx_gateway] if idx_gateway is not None else ""
                descr = r[idx_descr] if idx_descr is not None else ""
                drcr = r[idx_drcr] if idx_drcr is not None else ""

                # inherit section banner as gateway if empty
                if (gateway is None or str(gateway).strip() == "") and current_section:
                    gateway = current_section

                # label
                if force_type in {"Deposit", "Redeem"}:
                    label = force_type
                    label_src = f"force_type={force_type}"
                else:
                    label = _classify_label(
                        type_raw, status_raw, gateway, descr, drcr, sign
                    )
                    label_src = f"type/gateway/descr/drcr/sign"

                # --- force direction by status, if status is decisive ---
                from processor.normalization import normalize_status

                ns_bucket = normalize_status(
                    status_raw
                )  # -> 'accepted_withdraw' | 'accepted_deposit' | ...
                if ns_bucket == "accepted_withdraw":
                    label = "Redeem"
                elif ns_bucket == "accepted_deposit":
                    label = "Deposit"

                # Extra safety for raw tokens (some exports keep vendor tokens verbatim)
                s_raw = (str(status_raw) or "").lower()
                if any(
                    tok in s_raw
                    for tok in (
                        "rdm_closed",
                        "pay_closed",
                        "payout_closed",
                        "payout_paid",
                    )
                ):
                    label = "Redeem"

                # Extra safety: check normalized status for withdrawal indicators
                ns = normalize_status(status_raw)
                if ns in ("rdm_closed", "pay_closed", "payout_closed", "payout_paid"):
                    label = "Redeem"

                # --- STATUS FILTERS ---
                if had_status_col:
                    # compact A/C/P/W handling
                    s_up = str(status_raw or "").strip().upper()
                    if s_up in STATUS_ACPW:
                        if s_up != "A":
                            t(
                                f"row {r_i}: ⏭️ drop — compact status={s_up} (keep only A)"
                            )
                            continue
                        norm_status = STATUS_ACPW["A"]
                    else:
                        norm_status = _status_norm(status_raw)
                        if _is_bad_status(norm_status):
                            t(
                                f"row {r_i}: ⏭️ drop — bad status {norm_status!r} (type_raw={type_raw!r})"
                            )
                            continue
                        if not _is_final_success(norm_status):
                            # Get normalized status for better diagnostics
                            ns = normalize_status(status_raw)
                            t(
                                f"row {r_i}: ⏭️ drop — not final '{status_raw}' (normalized='{ns}') (type_raw={type_raw!r})"
                            )
                            continue
                else:
                    norm_status = "auto"

                rows.append([currency, dt, label, float(amount_val), norm_status])
                t(
                    f"row {r_i}: ✅ accept → dt={dt:%Y-%m-%d} amt={amount_val:.2f} sign={sign:+d} "
                    f"curr={currency} label={label} ({label_src}) status={norm_status!r}"
                )

            except Exception as e:
                t(f"row {r_i}: ⚠️ error — {e!r}")
                continue

        cur = end + 1

    dep = sum(r[3] for r in rows if r[2] == "Deposit")
    wd = sum(r[3] for r in rows if r[2] == "Redeem")
    net = dep - wd
    t(f"recap: deposits={dep:.2f}, cashouts={wd:.2f}, net={net:.2f}")
    return rows


def _write_payment_summary_sheet(excel_file, rows, log):
    """Internal helper: creates/replaces 'Payment Summary' from parsed rows."""
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = excel_file.wb
    if "Payment Summary" in wb.sheetnames:
        del wb["Payment Summary"]
    sh = wb.create_sheet("Payment Summary")

    headers = ["Date", "Type", "Amount", "Status"]
    for c, h in enumerate(headers, start=1):
        sh.cell(row=1, column=c).value = h

    for i, row in enumerate(rows, start=2):
        for j, v in enumerate(row, start=1):
            sh.cell(row=i, column=j).value = v

    end_col = get_column_letter(len(headers))
    end_row = len(rows) + 1
    ref = f"A1:{end_col}{end_row}"
    tbl = Table(displayName="PaymentSummaryTable", ref=ref)
    tbl.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9", showRowStripes=True, showColumnStripes=True
    )
    sh.add_table(tbl)

    dep = sum(r[2] for r in rows if r[1] == "Deposit")
    wd = sum(r[2] for r in rows if r[1] == "Redeem")
    log(f"TA logic: Deposits = {dep}, Cashouts = {abs(wd)}, Net = {dep + wd}")
    log(f"✅ Created Payment Summary with {len(rows)} rows.")


def create_ta_yearly_summary(excel_file, log):
    if "Payment Summary" not in excel_file.wb.sheetnames:
        log("⚠️ No Payment Summary sheet found.")
        return

    payment_sheet = excel_file.wb["Payment Summary"]
    if "Summary" in excel_file.wb.sheetnames:
        del excel_file.wb["Summary"]
    summary_sheet = excel_file.wb.create_sheet("Summary")

    summary_sheet["A1"] = "Year"
    summary_sheet["B1"] = "Total Deposits"
    summary_sheet["C1"] = "Total Cashouts"
    summary_sheet["D1"] = "Net"

    yearly = defaultdict(lambda: {"Deposit": 0.0, "Redeem": 0.0})
    for row in payment_sheet.iter_rows(min_row=2, values_only=True):
        dt, tx_type, amount, *_ = row  # Tolerate extra columns
        if not isinstance(dt, datetime):
            continue
        year = dt.year
        yearly[year][tx_type] += amount if tx_type == "Deposit" else -amount

    if not yearly:
        log("⚠️ No valid dates found — skipping Summary sheet creation.")
        return

    for i, year in enumerate(sorted(yearly.keys()), start=2):
        deposits = yearly[year]["Deposit"]
        cashouts = yearly[year]["Redeem"]
        net = deposits - cashouts
        summary_sheet.cell(row=i, column=1).value = year
        summary_sheet.cell(row=i, column=2).value = deposits
        summary_sheet.cell(row=i, column=3).value = cashouts
        summary_sheet.cell(row=i, column=4).value = net

    format_summary_sheet(summary_sheet)
    log("✅ Created Yearly Summary table.")


def format_summary_sheet(sheet):
    bold_font = Font(bold=True)
    center = Alignment(horizontal="center")
    border = Border(bottom=Side(style="thin"))
    for col in sheet.iter_cols(min_row=1, max_row=1):
        for cell in col:
            cell.font = bold_font
            cell.alignment = center
            cell.border = border
            cell.fill = PatternFill(
                start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"
            )

    for col in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        sheet.column_dimensions[get_column_letter(col[0].column)].width = max_length + 2


def create_ta_yearly_summary(excel_file, log):
    if "Payment Summary" not in excel_file.wb.sheetnames:
        log("⚠️ No Payment Summary sheet found.")
        return

    payment_sheet = excel_file.wb["Payment Summary"]
    rows = list(payment_sheet.iter_rows(min_row=2, values_only=True))

    # Aggregate data by year and type
    yearly_data = defaultdict(lambda: defaultdict(float))
    years = set()

    for row in rows:
        dt, tx_type, amount, *_ = row  # Tolerate extra columns
        if not isinstance(dt, datetime):
            continue
        year = dt.year
        years.add(year)
        yearly_data[tx_type][year] += amount

    sorted_years = sorted(years)
    tx_types = list(yearly_data.keys())

    # Create or clear Summary sheet
    if "Summary" in excel_file.wb.sheetnames:
        del excel_file.wb["Summary"]
    summary = excel_file.wb.create_sheet("Summary")

    # Title
    summary.merge_cells(
        start_row=1, start_column=1, end_row=1, end_column=len(sorted_years) + 2
    )
    summary["A1"] = f"Summary Payment Transactions ({min(years)} - {max(years)})"
    summary["A1"].font = Font(size=14, bold=True)
    summary["A1"].alignment = Alignment(horizontal="center")

    # Header row
    summary["A2"] = "Transaction Type"
    for i, year in enumerate(sorted_years):
        summary.cell(row=2, column=i + 2, value=year)
    summary.cell(row=2, column=len(sorted_years) + 2, value="Grand Total")

    # Fill in data
    for row_idx, tx_type in enumerate(tx_types, start=3):
        summary.cell(row=row_idx, column=1, value=tx_type)
        row_total = 0
        for col_idx, year in enumerate(sorted_years, start=2):
            amount = yearly_data[tx_type][year]
            summary.cell(row=row_idx, column=col_idx, value=round(amount, 2))
            row_total += amount
        summary.cell(
            row=row_idx, column=len(sorted_years) + 2, value=round(row_total, 2)
        )

    # Grand total row
    grand_row = len(tx_types) + 3
    summary.cell(row=grand_row, column=1, value="Grand Total")
    for col in range(2, len(sorted_years) + 3):
        col_letter = get_column_letter(col)
        formula = f"=SUM({col_letter}3:{col_letter}{grand_row-1})"
        summary.cell(row=grand_row, column=col, value=formula)

    # Table styling
    from processor.utils import format_summary_sheet

    format_summary_sheet(
        summary, header_row=2, last_row=grand_row, last_col=len(sorted_years) + 2
    )

    table_range = f"A2:{get_column_letter(len(sorted_years)+2)}{grand_row}"
    table = Table(displayName="SummaryTable", ref=table_range)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    summary.add_table(table)

    log("✅ Created Yearly Summary table.")


def format_summary_sheet(sheet, header_row, total_row, max_col):
    for row in sheet.iter_rows(
        min_row=header_row, max_row=total_row, min_col=1, max_col=max_col
    ):
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin"),
            )
            if cell.row == header_row:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(
                    start_color="4F81BD", end_color="4F81BD", fill_type="solid"
                )
            elif cell.column == 1:
                cell.font = Font(bold=True)
                cell.fill = PatternFill(
                    start_color="D9EAD3", end_color="D9EAD3", fill_type="solid"
                )
            elif cell.row == total_row:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(
                    start_color="C0504D", end_color="C0504D", fill_type="solid"
                )


def create_currency_type_date_summary(excel_file, log):
    # Replace with pivotish summary: Currency → Type → Year
    create_pivotish_summary_from_payment_summary(excel_file, log)


# --- New: Pivot-ish summary (Currency → Type → Year)
INDENT1 = Alignment(horizontal="left", indent=1)
INDENT2 = Alignment(horizontal="left", indent=2)
CENTER = Alignment(horizontal="center")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=13)
THIN = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

CURR_FILL = PatternFill("solid", fgColor="F2F2F2")
TYPE_FILL = PatternFill("solid", fgColor="FAFAFA")
TOT_FILL = PatternFill("solid", fgColor="FFF2CC")


def _auto_width(ws):
    for col in ws.columns:
        w = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = w + 2


def _canon_type(raw):
    s = (str(raw or "")).lower()
    return (
        "Withdrawal"
        if any(k in s for k in ("withdraw", "cashout", "cash out", "redeem", "payout"))
        else "Deposit"
    )


def _cube(rows):
    # cube[currency][type][year] = total
    cube = defaultdict(lambda: defaultdict(lambda: defaultdict(float)))
    for r in rows:
        cur = (r.get("currency") or "").strip() or "—"
        typ = _canon_type(r.get("type"))
        amt = float(r.get("amount") or 0)
        # force sign convention
        amt = -abs(amt) if typ == "Withdrawal" else abs(amt)
        if not is_success_status((r.get("status") or "").lower()):
            continue
        dt = r.get("date")
        if isinstance(dt, str):
            dt = parse_any_datetime(dt) or dt
        year = getattr(dt, "year", None)
        if year is None:
            continue
        cube[cur][typ][year] += amt
    return cube


def _gather_rows_for_summary(rows):
    """
    Normalize rows and collect cube plus min/max date for successful transactions.
    """
    cube = defaultdict(lambda: defaultdict(lambda: defaultdict(float)))
    min_dt = None
    max_dt = None

    for r in rows:
        status = (r.get("status") or "").lower()
        if not is_success_status(status):
            continue

        cur = (r.get("currency") or "").strip() or "—"
        typ = _canon_type(r.get("type"))
        amt = float(r.get("amount") or 0)
        amt = -abs(amt) if typ == "Withdrawal" else abs(amt)

        dt = r.get("date")
        if isinstance(dt, str) or dt is None:
            dt = parse_any_datetime(dt) if dt else None
        if dt is None:
            continue

        cube[cur][typ][dt.year] += amt
        if min_dt is None or dt < min_dt:
            min_dt = dt
        if max_dt is None or dt > max_dt:
            max_dt = dt

    return cube, min_dt, max_dt


def _rows_from_payment_summary(ws):
    headers = [str(c.value or "").strip().lower() for c in ws[1]]
    idx = {}
    for i, h in enumerate(headers):
        if "curr" in h:
            idx["currency"] = i
        elif "date" in h or "created" in h or "start" in h:
            idx["date"] = i
        elif "type" in h:
            idx["type"] = i
        elif "amount" in h:
            idx["amount"] = i
        elif "status" in h:
            idx["status"] = i

    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        out.append(
            {
                "currency": (
                    row[idx.get("currency")]
                    if idx.get("currency") is not None
                    else None
                ),
                "date": row[idx.get("date")] if idx.get("date") is not None else None,
                "type": row[idx.get("type")] if idx.get("type") is not None else None,
                "amount": (
                    row[idx.get("amount")] if idx.get("amount") is not None else None
                ),
                "status": (
                    row[idx.get("status")] if idx.get("status") is not None else None
                ),
            }
        )
    return out


def create_pivotish_summary_from_payment_summary(excel_file, log):
    wb = excel_file.wb
    if "Payment Summary" not in wb.sheetnames:
        log("ℹ️ No 'Payment Summary' sheet found. Skipping Summary.")
        return

    rows = _rows_from_payment_summary(wb["Payment Summary"])
    cube, start_dt, end_dt = _gather_rows_for_summary(rows)

    if "Summary" in wb.sheetnames:
        del wb["Summary"]
    ws = wb.create_sheet("Summary")

    # Title row (row 1)
    if start_dt and end_dt:
        title = (
            f"Summary of payment transactions ({start_dt:%d.%m.%Y} - {end_dt:%d.%m.%Y})"
        )
    else:
        title = "Summary of payment transactions"
    ws.merge_cells("A1:B1")
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws["A1"].alignment = CENTER

    # Table headers (row 2)
    ws["A2"] = "Row Labels"
    ws["B2"] = "Sum of Amount"
    ws["A2"].font = BOLD
    ws["B2"].font = BOLD

    r = 3
    grand = 0.0

    for cur in sorted(cube.keys()):
        cur_total = sum(sum(years.values()) for years in cube[cur].values())

        a = ws.cell(row=r, column=1, value=cur)
        a.font = BOLD
        a.fill = CURR_FILL
        c = ws.cell(row=r, column=2, value=cur_total)
        excel_fmt_number(c)
        c.alignment = Alignment(horizontal="right")
        c.fill = CURR_FILL
        grand += cur_total
        r += 1

        for tx in ("Deposit", "Withdrawal"):
            if tx not in cube[cur]:
                continue
            type_total = sum(cube[cur][tx].values())

            a = ws.cell(row=r, column=1, value=tx)
            a.font = BOLD
            a.alignment = INDENT1
            a.fill = TYPE_FILL
            c = ws.cell(row=r, column=2, value=type_total)
            excel_fmt_number(c)
            c.alignment = Alignment(horizontal="right")
            c.fill = TYPE_FILL
            r += 1

            for y in sorted(cube[cur][tx].keys()):
                a = ws.cell(row=r, column=1, value=y)
                a.alignment = INDENT2
                c = ws.cell(row=r, column=2, value=cube[cur][tx][y])
                excel_fmt_number(c)
                c.alignment = Alignment(horizontal="right")
                r += 1

        r += 1

    a = ws.cell(row=r, column=1, value="Grand Total")
    a.font = BOLD
    a.fill = TOT_FILL
    c = ws.cell(row=r, column=2, value=grand)
    excel_fmt_number(c)
    c.alignment = Alignment(horizontal="right")
    c.fill = TOT_FILL

    # Borders for the table area (headers + data)
    for row in ws.iter_rows(min_row=2, max_row=r, min_col=1, max_col=2):
        for cell in row:
            cell.border = THIN

    _auto_width(ws)
    ws.freeze_panes = "A3"
    log("✅ Created styled Summary (Currency → Type → Year) with title row.")


import re
from datetime import datetime
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

_NORM = re.compile(r"[^a-z0-9]+")


def _norm(s):
    if s is None:
        return ""
    return _NORM.sub("", str(s).strip().lower())


def _find_header_row(sheet, log, scan_rows=2500):
    """
    Find the header row by scanning the first `scan_rows` rows and looking
    for common payment column names (robust to case/spacing).
    Returns: (row_index_1_based, raw_header_values) or (None, None)
    """
    import re

    def norm(s):
        if s is None:
            return ""
        s = str(s).strip().lower()
        # collapse to letters/digits so "Transaction Type" -> "transactiontype"
        return re.sub(r"[^a-z0-9]+", "", s)

    # What we consider a "header"
    MARKERS = {
        "type": {
            "type",
            "transtype",
            "activity",
            "operation",
            "transactiontype",
            "transactiontypename",
        },
        "amount": {"amount", "txnamount", "initial", "net", "value", "billingamount"},
        "currency": {
            "currency",
            "txncurrency",
            "currencyid",
            "curr",
            "accountcurrency",
            "billingcurrency",
        },
        "status": {"status", "state", "result"},
        "date": {"date", "created", "started", "timestamp", "time", "datetime"},
        # optional extras, not required but strengthen the match
        "id": {"transid", "transactionid", "parentid", "reference", "refid"},
        "gateway": {"gateway", "processor", "method", "channel"},
    }

    def row_hits(normed_cells):
        hits = set()
        for cell in normed_cells:
            for group, names in MARKERS.items():
                if cell in names:
                    hits.add(group)
        return hits

    # Scan rows until we find one that looks like a header.
    for r_idx, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=scan_rows, values_only=True), start=1
    ):
        vals = [c for c in row]
        if not any(v not in (None, "", " ") for v in vals):
            continue

        nvals = [norm(v) for v in vals]
        hits = row_hits(nvals)

        # Heuristics:
        # - must include at least 3 core groups
        # - and must include either type or amount (to avoid date-only rows)
        core = {"type", "amount", "currency", "status", "date"}
        if len(hits & core) >= 3 and ({"type", "amount"} & hits):
            log(f"✅ Detected header row at Excel row {r_idx}: {vals}")
            return r_idx, list(vals)

        # Special-case: exact PokerStars audit style
        if {"transid", "transactiontype", "currency", "amount"} <= set(nvals):
            log(f"✅ Detected header row at Excel row {r_idx}: {vals}")
            return r_idx, list(vals)

    # Not found
    return None, None


def _parse_amount(x):
    if x is None:
        return None
    s = str(x).strip()
    # remove thousands separators, normalize decimal
    s = s.replace(" ", "").replace("'", "")
    # If both separators present, assume European (1.234,56)
    if "." in s and "," in s:
        s = s.replace(".", "").replace(",", ".")
    else:
        # If only comma present, treat as decimal separator
        if "," in s and "." not in s:
            s = s.replace(",", ".")
    try:
        return float(s)
    except Exception:
        return None
